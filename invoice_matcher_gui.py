#!/usr/bin/env python3
import json
import os
import queue
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
from tkinter import font as tkfont

import invoice_matcher as matcher


matcher.RAISE_ON_ERROR = True

APP_TITLE = "发票/订单匹配助手"
APP_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG_PATH = APP_DIR / "invoice_matcher_config.json"
DEFAULT_OUTPUT_PATH = APP_DIR / "match_result.xlsx"
DEFAULT_ORDER_FILES = ["订单数据.xlsx", "订单数据 (1).xlsx"]

PREFERRED_COLUMNS = {
    "order_id": ["订单号", "订单编号", "订单ID"],
    "amount": ["实付金额", "商品金额", "订单金额", "总金额", "金额"],
    "date": ["订单提交时间", "下单时间", "支付时间", "订单时间"],
    "vendor": ["店铺名称", "店铺", "商家", "卖家"],
    "description": ["商品名称", "商品", "明细", "描述"],
}


def read_order_columns(path):
    try:
        import pandas as pd

        df = pd.read_excel(path, nrows=1)
        return [str(c).strip() for c in df.columns]
    except Exception:
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise RuntimeError(
                "无法读取 Excel，请安装 pandas 或 openpyxl。"
            ) from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        wb.close()
        if not rows:
            return []
        return [str(c).strip() if c is not None else "" for c in rows[0]]


def auto_pick_column(columns, preferred_list):
    for prefer in preferred_list:
        for col in columns:
            if col == prefer or prefer in col:
                return col
    return ""


def parse_float(text, fallback):
    text = str(text).strip()
    if not text:
        return fallback
    try:
        return float(text)
    except ValueError:
        return fallback


def parse_int(text, fallback):
    text = str(text).strip()
    if not text:
        return fallback
    try:
        return int(text)
    except ValueError:
        return fallback


class InvoiceMatcherGUI:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.minsize(960, 680)

        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
        else:
            style.theme_use("clam")

        default_font = tkfont.nametofont("TkDefaultFont")
        default_font.configure(family="Microsoft YaHei UI", size=10)
        self.root.option_add("*Font", default_font)

        self.queue = queue.Queue()
        self.worker_thread = None

        self.orders_var = tk.StringVar(value=self._default_orders_path())
        self.pdf_dir_var = tk.StringVar(value=self._default_pdf_dir())
        self.config_var = tk.StringVar(value=str(DEFAULT_CONFIG_PATH))
        self.output_var = tk.StringVar(value=str(DEFAULT_OUTPUT_PATH))
        self.recursive_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="就绪")

        self.column_vars = {
            "order_id": tk.StringVar(),
            "amount": tk.StringVar(),
            "date": tk.StringVar(),
            "vendor": tk.StringVar(),
            "description": tk.StringVar(),
        }
        self.amount_tol_var = tk.StringVar(value="0.01")
        self.date_tol_var = tk.StringVar(value="7")
        self.min_score_var = tk.StringVar(value="0.7")
        self.allow_dupe_var = tk.BooleanVar(value=False)

        self._build_ui()
        self._load_config_to_ui()

    def _default_orders_path(self):
        for name in DEFAULT_ORDER_FILES:
            path = APP_DIR / name
            if path.exists():
                return str(path)
        return ""

    def _default_pdf_dir(self):
        candidate = APP_DIR / "发票PDF"
        return str(candidate) if candidate.exists() else ""

    def _build_ui(self):
        notebook = ttk.Notebook(self.root)
        notebook.pack(fill="both", expand=True)

        self.match_tab = ttk.Frame(notebook)
        self.config_tab = ttk.Frame(notebook)
        notebook.add(self.match_tab, text="匹配")
        notebook.add(self.config_tab, text="配置")

        self._build_match_tab()
        self._build_config_tab()

    def _build_match_tab(self):
        self.match_tab.columnconfigure(0, weight=1)
        self.match_tab.rowconfigure(2, weight=1)

        file_frame = ttk.LabelFrame(self.match_tab, text="文件与输出")
        file_frame.grid(row=0, column=0, sticky="nsew", padx=12, pady=10)
        file_frame.columnconfigure(1, weight=1)

        self._add_path_row(
            file_frame, 0, "订单文件 (.xlsx)", self.orders_var, self._choose_orders_file
        )
        self._add_path_row(
            file_frame, 1, "发票 PDF 文件夹", self.pdf_dir_var, self._choose_pdf_dir
        )
        self._add_path_row(
            file_frame, 2, "配置文件 (.json)", self.config_var, self._choose_config_file
        )
        self._add_path_row(
            file_frame, 3, "输出结果", self.output_var, self._choose_output_file
        )

        options_frame = ttk.Frame(file_frame)
        options_frame.grid(row=4, column=0, columnspan=3, sticky="w", padx=6, pady=(4, 8))
        ttk.Checkbutton(
            options_frame, text="递归扫描子文件夹", variable=self.recursive_var
        ).pack(anchor="w")

        action_frame = ttk.Frame(self.match_tab)
        action_frame.grid(row=1, column=0, sticky="ew", padx=12)
        action_frame.columnconfigure(0, weight=1)

        self.run_button = ttk.Button(
            action_frame, text="开始匹配", command=self._start_match
        )
        self.run_button.pack(side="left")

        self.open_output_button = ttk.Button(
            action_frame, text="打开结果", command=self._open_output
        )
        self.open_output_button.pack(side="left", padx=(10, 0))
        self.open_output_button.configure(state="disabled")

        self.progress = ttk.Progressbar(action_frame, mode="determinate")
        self.progress.pack(side="right", fill="x", expand=True, padx=(10, 0))

        log_frame = ttk.LabelFrame(self.match_tab, text="运行日志")
        log_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=10)
        log_frame.rowconfigure(0, weight=1)
        log_frame.columnconfigure(0, weight=1)

        self.log_text = ScrolledText(log_frame, height=12, wrap="word")
        self.log_text.grid(row=0, column=0, sticky="nsew")
        self.log_text.configure(state="disabled")

        status_frame = ttk.Frame(self.match_tab)
        status_frame.grid(row=3, column=0, sticky="ew", padx=12, pady=(0, 10))
        status_frame.columnconfigure(0, weight=1)
        ttk.Label(status_frame, textvariable=self.status_var).grid(row=0, column=0, sticky="w")

    def _build_config_tab(self):
        self.config_tab.columnconfigure(0, weight=1)
        self.config_tab.rowconfigure(2, weight=1)

        config_path_frame = ttk.LabelFrame(self.config_tab, text="配置文件")
        config_path_frame.grid(row=0, column=0, sticky="ew", padx=12, pady=10)
        config_path_frame.columnconfigure(1, weight=1)
        self._add_path_row(
            config_path_frame,
            0,
            "配置文件路径",
            self.config_var,
            self._choose_config_file,
        )

        mapping_frame = ttk.LabelFrame(self.config_tab, text="订单列映射")
        mapping_frame.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 10))
        mapping_frame.columnconfigure(1, weight=1)

        self._add_combo_row(
            mapping_frame, 0, "订单号", self.column_vars["order_id"]
        )
        self._add_combo_row(
            mapping_frame, 1, "实付金额", self.column_vars["amount"]
        )
        self._add_combo_row(
            mapping_frame, 2, "订单时间", self.column_vars["date"]
        )
        self._add_combo_row(
            mapping_frame, 3, "店铺名称", self.column_vars["vendor"]
        )
        self._add_combo_row(
            mapping_frame, 4, "商品名称", self.column_vars["description"]
        )

        read_columns_button = ttk.Button(
            mapping_frame, text="读取订单列名", command=self._load_columns_from_orders
        )
        read_columns_button.grid(row=5, column=0, pady=(6, 8), sticky="w")

        rules_frame = ttk.LabelFrame(self.config_tab, text="匹配规则")
        rules_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=(0, 10))
        rules_frame.columnconfigure(1, weight=1)

        ttk.Label(rules_frame, text="金额容差").grid(
            row=0, column=0, sticky="w", padx=6, pady=4
        )
        ttk.Entry(rules_frame, textvariable=self.amount_tol_var).grid(
            row=0, column=1, sticky="ew", padx=6, pady=4
        )
        ttk.Label(rules_frame, text="日期容差(天)").grid(
            row=1, column=0, sticky="w", padx=6, pady=4
        )
        ttk.Entry(rules_frame, textvariable=self.date_tol_var).grid(
            row=1, column=1, sticky="ew", padx=6, pady=4
        )
        ttk.Label(rules_frame, text="最低匹配分").grid(
            row=2, column=0, sticky="w", padx=6, pady=4
        )
        ttk.Entry(rules_frame, textvariable=self.min_score_var).grid(
            row=2, column=1, sticky="ew", padx=6, pady=4
        )
        ttk.Checkbutton(
            rules_frame,
            text="允许一个订单匹配多张发票",
            variable=self.allow_dupe_var,
        ).grid(row=3, column=0, columnspan=2, sticky="w", padx=6, pady=(4, 8))

        save_button = ttk.Button(
            rules_frame, text="保存配置", command=self._save_config
        )
        save_button.grid(row=4, column=0, columnspan=2, sticky="w", padx=6, pady=(4, 8))

    def _add_path_row(self, parent, row, label, variable, command):
        ttk.Label(parent, text=label).grid(
            row=row, column=0, sticky="w", padx=6, pady=6
        )
        entry = ttk.Entry(parent, textvariable=variable)
        entry.grid(row=row, column=1, sticky="ew", padx=6, pady=6)
        ttk.Button(parent, text="选择...", command=command).grid(
            row=row, column=2, sticky="ew", padx=6, pady=6
        )

    def _add_combo_row(self, parent, row, label, variable):
        ttk.Label(parent, text=label).grid(
            row=row, column=0, sticky="w", padx=6, pady=4
        )
        combo = ttk.Combobox(parent, textvariable=variable, values=[])
        combo.grid(row=row, column=1, sticky="ew", padx=6, pady=4)
        combo.configure(state="normal")

    def _log(self, message):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", message + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _choose_orders_file(self):
        path = filedialog.askopenfilename(
            title="选择订单文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.orders_var.set(path)

    def _choose_pdf_dir(self):
        path = filedialog.askdirectory(title="选择发票 PDF 文件夹")
        if path:
            self.pdf_dir_var.set(path)

    def _choose_config_file(self):
        path = filedialog.askopenfilename(
            title="选择配置文件",
            filetypes=[("配置文件", "*.json"), ("所有文件", "*.*")],
        )
        if path:
            self.config_var.set(path)
            self._load_config_to_ui()

    def _choose_output_file(self):
        path = filedialog.asksaveasfilename(
            title="选择输出文件",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx"), ("CSV 文件", "*.csv"), ("所有文件", "*.*")],
        )
        if path:
            self.output_var.set(path)

    def _load_columns_from_orders(self):
        orders_path = self.orders_var.get().strip()
        if not orders_path:
            messagebox.showerror("缺少文件", "请先选择订单文件。")
            return
        if not Path(orders_path).exists():
            messagebox.showerror("文件不存在", "订单文件路径无效。")
            return
        try:
            columns = read_order_columns(orders_path)
        except Exception as exc:
            messagebox.showerror("读取失败", str(exc))
            return
        if not columns:
            messagebox.showerror("读取失败", "无法读取订单表头。")
            return
        self._update_combo_values(columns)
        self._auto_fill_columns(columns)
        messagebox.showinfo("已读取", "订单列名已加载。")

    def _update_combo_values(self, columns):
        for child in self.config_tab.winfo_children():
            if isinstance(child, ttk.LabelFrame) and child.cget("text") == "订单列映射":
                for combo in child.winfo_children():
                    if isinstance(combo, ttk.Combobox):
                        combo.configure(values=columns)

    def _auto_fill_columns(self, columns):
        for key, prefer in PREFERRED_COLUMNS.items():
            if self.column_vars[key].get():
                continue
            picked = auto_pick_column(columns, prefer)
            if picked:
                self.column_vars[key].set(picked)

    def _load_config_to_ui(self):
        config_path = Path(self.config_var.get().strip())
        if config_path.exists():
            try:
                with open(config_path, "r", encoding="utf-8") as fh:
                    config = json.load(fh)
            except Exception:
                config = matcher.DEFAULT_CONFIG
        else:
            config = matcher.DEFAULT_CONFIG

        order_cols = config.get("order_columns", {})
        self.column_vars["order_id"].set(order_cols.get("order_id", ""))
        self.column_vars["amount"].set(order_cols.get("amount", ""))
        self.column_vars["date"].set(order_cols.get("date", ""))
        self.column_vars["vendor"].set(order_cols.get("vendor", ""))
        self.column_vars["description"].set(order_cols.get("description", ""))

        rules = config.get("match_rules", {})
        self.amount_tol_var.set(str(rules.get("amount_tolerance", 0.01)))
        self.date_tol_var.set(str(rules.get("date_days_tolerance", 7)))
        self.min_score_var.set(str(rules.get("min_score", 0.7)))
        self.allow_dupe_var.set(bool(rules.get("allow_duplicate_orders", False)))

    def _save_config(self):
        config_path = Path(self.config_var.get().strip())
        if not config_path:
            messagebox.showerror("缺少路径", "请指定配置文件路径。")
            return
        if config_path.exists():
            should = messagebox.askyesno("覆盖确认", "配置文件已存在，是否覆盖？")
            if not should:
                return

        config = matcher.DEFAULT_CONFIG.copy()
        existing = {}
        if config_path.exists():
            try:
                with open(config_path, "r", encoding="utf-8") as fh:
                    existing = json.load(fh)
            except Exception:
                existing = {}

        if existing:
            config.update(existing)

        config["order_columns"] = {
            "order_id": self.column_vars["order_id"].get().strip(),
            "amount": self.column_vars["amount"].get().strip(),
            "date": self.column_vars["date"].get().strip(),
            "vendor": self.column_vars["vendor"].get().strip(),
            "description": self.column_vars["description"].get().strip(),
        }
        config.setdefault("match_rules", {})
        config["match_rules"]["amount_tolerance"] = parse_float(
            self.amount_tol_var.get(), 0.01
        )
        config["match_rules"]["date_days_tolerance"] = parse_int(
            self.date_tol_var.get(), 7
        )
        config["match_rules"]["min_score"] = parse_float(
            self.min_score_var.get(), 0.7
        )
        config["match_rules"]["allow_duplicate_orders"] = bool(
            self.allow_dupe_var.get()
        )

        try:
            with open(config_path, "w", encoding="utf-8") as fh:
                json.dump(config, fh, indent=2, ensure_ascii=False)
                fh.write("\n")
        except Exception as exc:
            messagebox.showerror("保存失败", f"无法保存配置：{exc}")
            return

        messagebox.showinfo("保存成功", "配置已保存。")

    def _start_match(self):
        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showinfo("正在运行", "匹配任务仍在进行中。")
            return

        orders_path = self.orders_var.get().strip()
        pdf_dir = self.pdf_dir_var.get().strip()
        config_path = self.config_var.get().strip()
        output_path = self.output_var.get().strip()

        if not orders_path or not Path(orders_path).exists():
            messagebox.showerror("缺少文件", "请选择有效的订单文件。")
            return
        if not pdf_dir or not Path(pdf_dir).exists():
            messagebox.showerror("缺少文件夹", "请选择有效的发票 PDF 文件夹。")
            return
        if not config_path:
            messagebox.showerror("缺少配置", "请指定配置文件。")
            return
        if not output_path:
            messagebox.showerror("缺少输出", "请指定输出文件。")
            return

        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

        self.progress.configure(value=0, maximum=1)
        self.status_var.set("开始处理...")
        self.run_button.configure(state="disabled")
        self.open_output_button.configure(state="disabled")

        self.worker_thread = threading.Thread(
            target=self._worker,
            args=(orders_path, pdf_dir, config_path, output_path, self.recursive_var.get()),
            daemon=True,
        )
        self.worker_thread.start()
        self.root.after(100, self._process_queue)

    def _worker(self, orders_path, pdf_dir, config_path, output_path, recursive):
        try:
            self.queue.put(("log", "读取配置文件..."))
            config = matcher.load_config(config_path)

            backend = matcher.resolve_pdf_backend()
            if backend is None:
                raise matcher.MatchError(
                    "缺少 PDF 解析库，请安装 pdfplumber 或 pypdf。"
                )
            self.queue.put(("log", f"PDF 解析后端: {backend}"))

            self.queue.put(("log", "读取订单数据..."))
            orders_df = matcher.load_orders(orders_path, config)
            self.queue.put(("log", f"订单行数: {len(orders_df)}"))

            self.queue.put(("log", "扫描发票 PDF..."))
            pdfs = matcher.find_pdfs(pdf_dir, recursive)
            if not pdfs:
                raise matcher.MatchError("未找到任何 PDF 文件。")
            self.queue.put(("progress_max", len(pdfs)))
            self.queue.put(("log", f"发票数量: {len(pdfs)}"))

            invoices = []
            for idx, path in enumerate(pdfs, 1):
                invoices.append(matcher.extract_invoice(path, config, backend))
                if idx % 5 == 0 or idx == len(pdfs):
                    self.queue.put(("progress", idx))

            self.queue.put(("log", "计算匹配结果..."))
            matches = matcher.match_orders_invoices(orders_df, invoices, config)
            matcher.write_output(output_path, orders_df, invoices, matches)

            self.queue.put(
                (
                    "done",
                    {
                        "orders": len(orders_df),
                        "invoices": len(invoices),
                        "matches": len(matches),
                        "output": output_path,
                    },
                )
            )
        except matcher.MatchError as exc:
            self.queue.put(("error", str(exc)))
        except Exception as exc:
            self.queue.put(("error", f"运行失败：{exc}"))

    def _process_queue(self):
        while True:
            try:
                kind, payload = self.queue.get_nowait()
            except queue.Empty:
                break

            if kind == "log":
                self._log(payload)
            elif kind == "progress_max":
                self.progress.configure(maximum=max(int(payload), 1))
            elif kind == "progress":
                self.progress.configure(value=int(payload))
                self.status_var.set(f"处理发票中：{payload}/{int(self.progress['maximum'])}")
            elif kind == "done":
                self._log(
                    f"完成：订单 {payload['orders']}，发票 {payload['invoices']}，匹配 {payload['matches']}。"
                )
                self.status_var.set("完成")
                if Path(payload["output"]).exists():
                    self.open_output_button.configure(state="normal")
            elif kind == "error":
                self._log(payload)
                self.status_var.set("出错")
                messagebox.showerror("匹配失败", payload)

        if self.worker_thread and self.worker_thread.is_alive():
            self.root.after(120, self._process_queue)
        else:
            self.run_button.configure(state="normal")
            if self.status_var.get() not in ("完成", "出错"):
                self.status_var.set("就绪")

    def _open_output(self):
        output_path = self.output_var.get().strip()
        if not output_path or not Path(output_path).exists():
            messagebox.showerror("文件不存在", "输出文件不存在。")
            return
        try:
            os.startfile(output_path)
        except Exception as exc:
            messagebox.showerror("无法打开", f"无法打开文件：{exc}")


def main():
    root = tk.Tk()
    app = InvoiceMatcherGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
